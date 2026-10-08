"""Facade: a buyer's spec workbook -> ``MeasurementChart`` (build plan M1, ``pom_import``).

Rows keep the legacy shape (``parser_version`` 'xlsx_v2') so the 1078983 goldens stay identical.
"""
import zipfile
from io import BytesIO

import openpyxl

from app.domain.measurement_chart import ChartRow, MeasurementChart
from app.domain.pom_labels import is_half
from app.domain.pom_mapping import map_row
from app.infrastructure.base_size import detect_base
from app.infrastructure.formulas import number
from app.infrastructure.header_meta import range_issues, read_meta
from app.infrastructure.pom_cells import read_cell
from app.infrastructure.pom_table import TableLayout, find_table

MAX_EXPANDED_BYTES = 30_000_000
MAX_ROWS, MAX_COLUMNS = 2000, 100


def _load(data: bytes):
    with zipfile.ZipFile(BytesIO(data)) as archive:
        if sum(info.file_size for info in archive.infolist()) > MAX_EXPANDED_BYTES:
            raise ValueError("Workbook expanded content exceeds 30 MB")
    book = openpyxl.load_workbook(BytesIO(data), data_only=False, keep_links=False)
    cached = openpyxl.load_workbook(BytesIO(data), data_only=True, keep_links=False)
    if any(sheet.max_row > MAX_ROWS or sheet.max_column > MAX_COLUMNS for sheet in book):
        book.close()
        cached.close()
        raise ValueError("Workbook exceeds 2000 rows or 100 columns")
    return book, cached


def _text(sheet, row: int, col: int | None) -> str:
    return str(sheet.cell(row, col).value or "").strip() if col else ""


def _tolerance(raw) -> float | None:
    try:
        return number(raw)
    except (ValueError, TypeError):
        return None


def _rows(sheet, cached_sheet, layout: TableLayout, filename: str):
    taken: set[str] = set()
    for row in range(layout.header_row + 1, sheet.max_row + 1):
        code, label = _text(sheet, row, layout.code_col), _text(sheet, row, layout.label_col)
        if not code or not label:
            continue
        alt = _text(sheet, row, layout.alt_col) or None
        values = {size: read_cell(sheet, cached_sheet, row, col) for size, col in layout.size_cols.items()}
        key, status, mapped_from = map_row(code, label, alt, taken)
        if status != "unmapped":
            taken.add(key)
        tolerance = sheet.cell(row, layout.tolerance_col).value if layout.tolerance_col else None
        known = sum(1 for cell in values.values() if cell["value"] is not None)
        yield {
            "key": key, "mapping_status": status, "mapped_from": mapped_from, "code": code,
            "label": label.lstrip("* "), "original_label": label, "parser_version": "xlsx_v2",
            "unit": "unconfirmed", "confidence": round(known / max(len(values), 1), 3),
            "tolerance": tolerance, "source": filename, "sheet": sheet.title, "row": row, "values": values,
        }, ChartRow(str(row), is_half(label), _tolerance(tolerance), alt)


def parse_chart(data: bytes, filename: str) -> MeasurementChart:
    book, cached = _load(data)
    try:
        layout = find_table(book)
        if layout is None:
            raise ValueError("No measurement rows with size headers found")
        sheet, cached_sheet = book[layout.sheet], cached[layout.sheet]
        pairs = list(_rows(sheet, cached_sheet, layout, filename))
        if not pairs:
            raise ValueError("No measurement rows with size headers found")
        sizes = tuple(layout.size_cols)
        meta = read_meta(cached_sheet, layout.header_row)
        return MeasurementChart(
            sheet=layout.sheet, header_row=layout.header_row, sizes=sizes,
            base_size=detect_base(sheet, layout.size_cols, [measurement["row"] for measurement, _ in pairs]),
            unit_hint=layout.unit_hint, rows=tuple(row for _, row in pairs),
            measurements=[measurement for measurement, _ in pairs], meta=meta, issues=range_issues(meta, sizes),
        )
    finally:
        book.close()
        cached.close()
