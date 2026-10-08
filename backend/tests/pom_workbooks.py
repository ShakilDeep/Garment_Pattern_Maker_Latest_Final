"""In-memory workbooks for the T126 POM import tests (nothing is written to disk)."""
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook

FIXTURES = Path(__file__).resolve().parents[2] / "fixtures"
CL1086464 = FIXTURES / "reference_patterns" / "CL1086464.xlsx"
BOOK2 = FIXTURES / "demo_sources" / "Book2(4).xlsx"


def workbook_bytes(*sheets: tuple[str, list[list]]) -> bytes:
    """Build a workbook from (title, rows) pairs; the first sheet replaces the default one."""
    book = Workbook()
    book.remove(book.active)
    for title, rows in sheets:
        sheet = book.create_sheet(title)
        for row in rows:
            sheet.append(row)
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def without_cached_values(data: bytes) -> bytes:
    """Re-save through openpyxl, which writes formulas but drops Excel's cached results."""
    buffer = BytesIO()
    load_workbook(BytesIO(data)).save(buffer)
    return buffer.getvalue()


def chart_values(chart, key: str) -> list:
    row = next(row for row in chart.measurements if row["key"] == key)
    return [row["values"][size]["value"] for size in chart.sizes]


def chart_row(chart, key: str):
    measurement = next(row for row in chart.measurements if row["key"] == key)
    return next(row for row in chart.rows if row.ref == str(measurement["row"]))
