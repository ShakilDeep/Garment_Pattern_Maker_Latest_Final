"""T125 (M0): the build-plan reference fixtures are present and byte-identical to the recorded sources."""

import hashlib
from io import BytesIO

import pytest
from conftest import REFERENCE_PATTERNS, needs_tech_packs
from openpyxl import load_workbook

RECORDED_SHA256 = {
    "1078375.PDS": "3bb70849a5567db409c55b4a74588a7b2074df78e54d55ea1eb14342538ba347",
    "1074664.PDS": "deae727f203f70a9570d2668a2c858edd7a8c27fa38bcf2ab1b3ea5bd7ab7e1d",
    "CL1086464.xlsx": "34c8e8e9a9311905567c007c0538cffc1a6dd96dc5a302f1f4ddc68253596993",
}
TECH_PACK_SHA256 = {
    "1078375.pdf": "ea2d3f7a37847bd861b09671f2d685965048e635966c1d192db6f11b637f48e5",
    "CL1086464.pdf": "03e6344dc8f17f0668e6a77bce96c063fb375b86c8b9b4608c24a5481b295ff2",
}


@pytest.mark.parametrize(("name", "digest"), RECORDED_SHA256.items())
def test_fixture_is_unchanged(name, digest):
    assert hashlib.sha256((REFERENCE_PATTERNS / name).read_bytes()).hexdigest() == digest


@needs_tech_packs
@pytest.mark.parametrize(("name", "digest"), TECH_PACK_SHA256.items())
def test_local_tech_pack_is_unchanged(name, digest):
    assert hashlib.sha256((REFERENCE_PATTERNS / name).read_bytes()).hexdigest() == digest


def test_readme_records_the_same_hashes():
    readme = (REFERENCE_PATTERNS / "README.md").read_text(encoding="utf-8")
    recorded = {**RECORDED_SHA256, **TECH_PACK_SHA256}
    assert all(f"`{name}`" in readme and digest in readme for name, digest in recorded.items())


def test_pant_measurement_workbook_opens():
    workbook = load_workbook(BytesIO((REFERENCE_PATTERNS / "CL1086464.xlsx").read_bytes()), read_only=True)
    assert workbook.sheetnames
