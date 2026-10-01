"""P0-08: every JSON API response is a typed model, and collection endpoints return Page[T]."""
from io import BytesIO
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from test_integrity import _ready

from app.api.main import create_app

UNTYPED = ("ObjectResponse", "ListResponse")
PAGED = ("/projects", "/projects/{pid}/measurements", "/projects/{pid}/documents",
         "/projects/{pid}/reviews", "/projects/{pid}/requirements/blocking")


def _json_schemas(spec):
    for path, operations in spec["paths"].items():
        for method, operation in operations.items():
            content = operation["responses"].get("200", {}).get("content", {})
            if "application/json" in content:
                yield path.removeprefix("/api/v1"), method, content["application/json"]["schema"]


@pytest.fixture(scope="module")
def spec(tmp_path_factory):
    return TestClient(create_app(f"sqlite:///{tmp_path_factory.mktemp('spec')}/spec.db")).get("/openapi.json").json()


def test_openapi_has_no_untyped_wrappers(spec):
    assert not set(UNTYPED) & set(spec["components"]["schemas"])


def test_every_json_response_names_a_model(spec):
    untyped = [(path, method) for path, method, schema in _json_schemas(spec)
               if "$ref" not in schema and "$ref" not in schema.get("items", {}) and path != "/health"]
    assert untyped == []


def test_collection_endpoints_are_paginated(spec):
    refs = {path: schema.get("$ref", "") for path, method, schema in _json_schemas(spec) if method == "get"}
    assert all(refs[path].split("/")[-1].startswith("Page_") for path in PAGED)


def test_project_list_pages_with_total(tmp_path):
    client = TestClient(create_app(f"sqlite:///{tmp_path}/pages.db"))
    for name in ("Alpha", "Beta", "Gamma"):
        client.post("/api/v1/projects", json={"name": name})
    page = client.get("/api/v1/projects", params={"limit": 2, "offset": 1}).json()
    assert (len(page["items"]), page["total"], page["limit"], page["offset"]) == (2, 3, 2, 1)
    assert client.get("/api/v1/projects", params={"offset": 5}).json()["items"] == []


@pytest.mark.parametrize("params", [{"limit": 0}, {"limit": 501}, {"offset": -1}])
def test_page_parameters_are_bounded(tmp_path, params):
    client = TestClient(create_app(f"sqlite:///{tmp_path}/bounds.db"))
    response = client.get("/api/v1/projects", params=params)
    assert response.status_code == 422
    assert response.json()["code"] == "REQUEST_INVALID"


def test_project_measurements_page_and_keep_undeclared_fields(tmp_path):
    client, url = _ready(tmp_path)
    page = client.get(f"{url}/measurements", params={"limit": 5}).json()
    assert len(page["items"]) == 5 and page["total"] > 5
    assert "mapping_status" in page["items"][0]
    assert "resolution_metadata" in client.get(url).json()


def _workbook_with_tolerance(text):
    book = load_workbook(Path(__file__).resolve().parents[2] / "references/Book2(4).xlsx")
    sheet = book.active
    header = next(row for row in sheet.iter_rows(max_row=30) if any("TOLERANCE" in str(c.value).upper() for c in row))
    column = next(c.column for c in header if "TOLERANCE" in str(c.value).upper())
    sheet.cell(header[0].row + 1, column).value = text
    data = BytesIO()
    book.save(data)
    return data.getvalue()


def test_text_tolerance_from_a_workbook_is_kept_verbatim(tmp_path):
    client = TestClient(create_app(f"sqlite:///{tmp_path}/tolerance.db"))
    url = "/api/v1/projects/" + client.post("/api/v1/projects", json={"name": "Tolerance"}).json()["id"]
    xlsx = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    files = {"file": ("chart.xlsx", _workbook_with_tolerance("±1/2"), xlsx)}
    assert client.post(f"{url}/documents", files=files).status_code == 200
    project = client.get(url)
    assert project.status_code == 200
    assert "±1/2" in [row["tolerance"] for row in project.json()["measurements"]]
